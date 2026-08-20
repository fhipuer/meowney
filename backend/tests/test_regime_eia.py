import io
from datetime import date, timedelta

import pytest
from openpyxl import Workbook

from app.config import settings
from app.services.regime_eia import (
    EiaPowerService,
    find_latest_eia860m_workbook,
    parse_eia860m_workbook,
    parse_eia_rows,
    parse_rto_daily_rows,
)
from app.services.regime_thesis import (
    aggregate_daily_series,
    aligned_daily_yoy,
    classify_grid_operations_axis,
    classify_interconnection_axis,
    classify_power_demand,
    classify_power_demand_axis,
    classify_power_investment_thesis,
    classify_power_supply_axis,
    classify_transmission_investment_axis,
    rolling_power_yoy_history,
    window_grid_operations,
)


def test_eia_parser_normalizes_monthly_and_annual_periods():
    monthly = parse_eia_rows(
        {"response": {"data": [{
            "period": "2026-05", "sectorid": "COM", "sales": "124043.5",
            "sales-units": "million kilowatt hours", "stateid": "US",
        }]}},
        {"COM": ("us_electricity_sales_commercial", "sales")}, "https://example.test",
    )
    annual = parse_eia_rows(
        {"response": {"data": [{
            "period": "2024", "stateID": "US", "net-summer-capacity": "1230416",
            "net-summer-capacity-units": "megawatts",
        }]}},
        {"US": ("us_electricity_net_summer_capacity", "net-summer-capacity")},
        "https://example.test",
    )

    assert monthly[0]["observation_date"] == "2026-05-01"
    assert monthly[0]["value"] == 124043.5
    assert annual[0]["observation_date"] == "2024-12-31"


def test_power_classification_uses_seasonally_matched_three_month_yoy():
    state, _ = classify_power_demand({"yoy_3m_avg": 4.0}, {"yoy_3m_avg": 5.0})
    assert state == "수요 확장"

    state, _ = classify_power_demand({"yoy_3m_avg": -1.0}, {"yoy_3m_avg": -2.0})
    assert state == "수요 둔화"

    state, _ = classify_power_demand({"yoy_3m_avg": 0.5}, {"yoy_3m_avg": 3.0})
    assert state == "상업용 수요 우세"


def test_rto_parser_uses_all_grid_lanes_and_one_common_day_boundary():
    payload = {"response": {"data": [
        {"period": "2026-08-19", "respondent": "US48", "type": "D",
         "timezone": "Eastern", "value": "1200", "value-units": "megawatthours"},
        {"period": "2026-08-19", "respondent": "TEX", "type": "D",
         "timezone": "Eastern", "value": "300", "value-units": "megawatthours"},
        {"period": "2026-08-19", "respondent": "TEX", "type": "D",
         "timezone": "Central", "value": "301", "value-units": "megawatthours"},
        {"period": "2026-08-19", "respondent": "TEX", "type": "DF",
         "timezone": "Eastern", "value": "310", "value-units": "megawatthours"},
        {"period": "2026-08-19", "respondent": "TEX", "type": "NG",
         "timezone": "Eastern", "value": "250", "value-units": "megawatthours"},
        {"period": "2026-08-19", "respondent": "TEX", "type": "TI",
         "timezone": "Eastern", "value": "-50", "value-units": "megawatthours"},
    ]}}

    rows = parse_rto_daily_rows(payload)

    assert [row["series_id"] for row in rows] == [
        "us_electricity_daily_demand_us48",
        "us_electricity_daily_demand_tex",
        "us_electricity_daily_demand_forecast_tex",
        "us_electricity_daily_net_generation_tex",
        "us_electricity_daily_total_interchange_tex",
    ]
    assert rows[1]["dimensions"]["timezone"] == "Eastern"


def test_latest_eia860m_workbook_is_selected_by_report_month():
    html = """
    <a href="/electricity/data/eia860m/xls/december_generator2025.xlsx">December</a>
    <a href="/electricity/data/eia860m/xls/june_generator2026.xlsx">June</a>
    """

    url, observed = find_latest_eia860m_workbook(html)

    assert url.endswith("/june_generator2026.xlsx")
    assert observed == "2026-06-30"


def _eia860m_fixture() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    operating = workbook.create_sheet("Operating")
    planned = workbook.create_sheet("Planned")
    operating.append([])
    operating.append([])
    operating.append([
        "Net Summer Capacity (MW)", "Status",
        "Planned Retirement Month", "Planned Retirement Year",
    ])
    operating.append([100, "(OP) Operating", 6, 2027])
    operating.append([1400, "(OP) Operating", None, None])
    operating.append([900, "(SB) Standby", 6, 2027])
    planned.append([])
    planned.append([])
    planned.append([
        "Net Summer Capacity (MW)", "Status", "Planned Operation Month",
        "Planned Operation Year", "Technology",
    ])
    planned.append([150, "(U) Under construction", 1, 2027, "Solar Photovoltaic"])
    planned.append([100, "(V) Under construction", 6, 2027, "Battery Storage"])
    planned.append([50, "(TS) Construction complete", 12, 2027, "Natural Gas Fired"])
    planned.append([999, "(T) Regulatory approvals received", 12, 2027, "Solar"])
    planned.append([50, "(U) Under construction", 12, 2029, "Wind"])
    planned.append([25, "(U) Under construction", 5, 2026, "Wind"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_eia860m_pipeline_counts_only_committed_near_term_projects():
    rows = parse_eia860m_workbook(
        _eia860m_fixture(),
        source_url="https://example.test/june_generator2026.xlsx",
        observation_date="2026-06-30",
    )
    values = {row["series_id"]: row["value"] for row in rows}

    assert values["us_power_operating_capacity_mw"] == 1500
    assert values["us_power_committed_additions_24m_mw"] == 300
    assert values["us_power_planned_retirements_24m_mw"] == 100
    assert values["us_power_net_committed_additions_24m_mw"] == 200
    assert values["us_power_delayed_committed_capacity_mw"] == 25
    assert values["us_power_committed_pipeline_solar_mw"] == 150
    assert values["us_power_committed_pipeline_battery_mw"] == 100
    assert values["us_power_committed_pipeline_gas_mw"] == 50
    assert values["us_power_net_pipeline_ratio_24m_pct"] == pytest.approx(200 / 15)
    assert values["us_power_variable_storage_share_24m_pct"] == pytest.approx(250 / 3)


def test_eia860m_does_not_turn_invalid_capacity_into_zero():
    workbook = Workbook()
    workbook.remove(workbook.active)
    operating = workbook.create_sheet("Operating")
    planned = workbook.create_sheet("Planned")
    for sheet, headers in (
        (operating, ["Net Summer Capacity (MW)", "Status", "Planned Retirement Month", "Planned Retirement Year"]),
        (planned, ["Net Summer Capacity (MW)", "Status", "Planned Operation Month", "Planned Operation Year", "Technology"]),
    ):
        sheet.append([])
        sheet.append([])
        sheet.append(headers)
    operating.append(["not-a-number", "(OP) Operating", None, None])
    operating.append([100, "(OP) Operating", None, None])
    planned.append(["missing", "(U) Under construction", 8, 2026, "Solar"])
    output = io.BytesIO()
    workbook.save(output)

    rows = parse_eia860m_workbook(
        output.getvalue(),
        source_url="https://example.test/eia860m.xlsx",
        observation_date="2026-06-30",
    )
    values = {row["series_id"]: row["value"] for row in rows}

    assert values["us_power_operating_capacity_mw"] == 100
    assert values["us_power_committed_additions_24m_mw"] == 0


def test_daily_yoy_uses_weekday_aligned_364_day_comparison():
    end = date(2026, 8, 19)
    points = []
    for offset in range(84):
        current = end - timedelta(days=offset)
        points.extend([
            {"date": current.isoformat(), "value": 110},
            {"date": (current - timedelta(days=364)).isoformat(), "value": 100},
        ])

    result = aligned_daily_yoy(points, window_days=84)

    assert result["value"] == pytest.approx(10)
    assert result["sample_days"] == 84


def test_power_history_plots_yoy_instead_of_seasonal_absolute_load():
    end = date(2026, 8, 19)
    national = []
    ai_regions = []
    for offset in range(60):
        current = end - timedelta(days=offset)
        prior = current - timedelta(days=364)
        national.extend([
            {"date": current.isoformat(), "value": 110},
            {"date": prior.isoformat(), "value": 100},
        ])
        ai_regions.extend([
            {"date": current.isoformat(), "value": 120},
            {"date": prior.isoformat(), "value": 100},
        ])

    history = rolling_power_yoy_history(national, ai_regions)

    assert history
    assert history[-1]["national"] == pytest.approx(10)
    assert history[-1]["ai_regions"] == pytest.approx(20)


@pytest.mark.parametrize(
    ("kwargs", "expected"),
    [
        ({"national_yoy_84d": 4, "ai_regions_yoy_84d": 5,
          "commercial_yoy_3m": 4, "regional_expansion_share": .8,
          "ai_regions_acceleration_pp": 1.2,
          "ai_excess_growth_pp": .5}, "전력 수요 빠르게 확대"),
        ({"national_yoy_84d": 2, "ai_regions_yoy_84d": 2,
          "commercial_yoy_3m": 1, "regional_expansion_share": .4,
          "ai_regions_acceleration_pp": 0,
          "ai_excess_growth_pp": 0}, "전력 수요 빠르게 확대"),
        ({"national_yoy_84d": -3, "ai_regions_yoy_84d": -3,
          "commercial_yoy_3m": -1, "regional_expansion_share": .2,
          "ai_regions_acceleration_pp": -1.2,
          "ai_excess_growth_pp": 0}, "전력 수요 감소"),
        ({"national_yoy_84d": 1, "ai_regions_yoy_84d": 3,
          "commercial_yoy_3m": 1, "regional_expansion_share": .5,
          "ai_regions_acceleration_pp": 0,
          "ai_excess_growth_pp": 2}, "AI 관찰지역 중심 확대"),
    ],
)
def test_power_demand_axis_has_deterministic_boundaries(kwargs, expected):
    state, _, _, coverage = classify_power_demand_axis(**kwargs)
    assert state == expected
    assert coverage == 1


def test_power_supply_and_composite_do_not_overclaim_grid_bottleneck():
    supply_state, _ = classify_power_supply_axis(4.9, 89)
    state, reason = classify_power_investment_thesis("전력 수요 확대", supply_state)

    assert supply_state == "건설 진행"
    assert state == "전력수요 확대·병목 미확인"
    assert "확정하지 않습니다" in reason


def test_fixed_region_basket_fails_closed_when_one_region_is_missing():
    complete = [[{"date": "2026-08-19", "value": value}] for value in (1, 2, 3, 4, 5)]

    assert aggregate_daily_series(complete, expected_series_count=6) == []
    assert aggregate_daily_series(complete, expected_series_count=5)[0]["value"] == 15


def test_eia_daily_cache_completeness_requires_every_type_and_region():
    service = EiaPowerService()
    required = {
        f"us_electricity_daily_{series_name}_{respondent.lower()}"
        for series_name in ("demand", "demand_forecast", "net_generation", "total_interchange")
        for respondent in ("us48", "cal", "car", "cent", "fla", "midw", "ne", "ny", "nw", "se", "sw", "ten", "tex", "mida")
    }

    class CacheRepo:
        def __init__(self, missing=None):
            self.missing = missing

        def series(self, series_id, *, limit=None):
            return [] if series_id == self.missing else ([{"value": 1}] if series_id in required else [])

    service.repo = CacheRepo()
    assert service._daily_grid_cache_complete() is True

    service.repo = CacheRepo(next(iter(required)))
    assert service._daily_grid_cache_complete() is False


def test_grid_operations_uses_matched_windows_and_negative_ti_as_imports():
    def points(value):
        return [
            {"date": (date(2026, 8, 19) - timedelta(days=index)).isoformat(), "value": value}
            for index in range(28)
        ]

    result = window_grid_operations(
        points(102), points(100), points(95), points(-7),
    )

    assert result["forecast_surprise_pct"] == pytest.approx(2)
    assert result["forecast_abs_error_pct"] == pytest.approx(2)
    assert result["generation_coverage_pct"] == pytest.approx(95 / 102 * 100)
    assert result["net_import_share_pct"] == pytest.approx(7 / 102 * 100)


def test_grid_operations_requires_load_plus_independent_pressure_context():
    state, _ = classify_grid_operations_axis(
        load_yoy_28d=3.5,
        forecast_surprise_pct=2,
        forecast_abs_error_pct=2.4,
        net_import_share_pct=6,
        pressure_region_count=4,
        expected_region_count=6,
        coverage=1,
    )
    assert state == "운영 부담 높음"

    state, _ = classify_grid_operations_axis(
        load_yoy_28d=0.2,
        forecast_surprise_pct=0.5,
        forecast_abs_error_pct=1,
        net_import_share_pct=1,
        pressure_region_count=0,
        expected_region_count=6,
        coverage=1,
    )
    assert state == "운영 여유"


@pytest.mark.parametrize(
    ("kwargs", "expected"),
    [
        ({
            "active_queue_gw": 2000, "operating_capacity_gw": 1300,
            "ia_executed_share_pct": 24.3, "median_active_age_years": 2.984,
            "ir_to_cod_median_years": 4.999, "usable": True,
        }, "접속 대기 부담 높음"),
        ({
            "active_queue_gw": 800, "operating_capacity_gw": 1300,
            "ia_executed_share_pct": 25, "median_active_age_years": 2,
            "ir_to_cod_median_years": 4, "usable": True,
        }, "접속 대기 부담"),
        ({
            "active_queue_gw": 500, "operating_capacity_gw": 1300,
            "ia_executed_share_pct": 70, "median_active_age_years": 1,
            "ir_to_cod_median_years": 2, "usable": True,
        }, "부담 완화"),
        ({
            "active_queue_gw": 2000, "operating_capacity_gw": 1300,
            "ia_executed_share_pct": 20, "median_active_age_years": 4,
            "ir_to_cod_median_years": 6, "usable": False,
        }, "자료 부족"),
    ],
)
def test_interconnection_axis_boundaries_and_freshness_gate(kwargs, expected):
    state, _ = classify_interconnection_axis(**kwargs)
    assert state == expected


@pytest.mark.parametrize(
    ("cagr", "reporters", "coverage", "usable", "expected"),
    [
        (8, 200, 90, True, "송전 투자 확대"),
        (0, 200, 90, True, "투자 유지"),
        (-0.1, 200, 90, True, "투자 둔화"),
        (12, 49, 90, True, "표본 제한"),
        (12, 200, 84.9, True, "표본 제한"),
        (12, 200, 90, False, "자료 부족"),
    ],
)
def test_transmission_axis_boundaries_and_freshness_gate(
    cagr, reporters, coverage, usable, expected,
):
    state, _ = classify_transmission_investment_axis(
        additions_usd=35_000_000_000,
        like_for_like_cagr_pct=cagr,
        reporter_count=reporters,
        reporter_coverage_pct=coverage,
        usable=usable,
    )
    assert state == expected


@pytest.mark.asyncio
async def test_eia_refresh_saves_all_required_series(monkeypatch):
    class FakeClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_):
            return None

    def row(series_id):
        return {
            "series_id": series_id, "observation_date": "2026-06-30",
            "value": 1, "unit": "MW", "source_url": "https://example.test",
        }

    async def monthly(_service, _client):
        return [row("us_electricity_sales_all")]

    async def daily(_service, _client):
        return [row("us_electricity_daily_demand_us48")]

    async def pipeline(_service, _client):
        return [row("us_power_operating_capacity_mw")]

    monkeypatch.setattr(settings, "eia_api_key", "test-key")
    monkeypatch.setattr("app.services.regime_eia.httpx.AsyncClient", lambda **_: FakeClient())
    monkeypatch.setattr(EiaPowerService, "_fetch_monthly_context", monthly)
    monkeypatch.setattr(EiaPowerService, "_fetch_daily_grid", daily)
    monkeypatch.setattr(EiaPowerService, "_fetch_capacity_pipeline", pipeline)

    result = await EiaPowerService().refresh(force=True)

    assert result == {"status": "success", "saved": 3}


@pytest.mark.asyncio
async def test_eia_refresh_keeps_cached_pipeline_but_reports_partial_failure(monkeypatch):
    class FakeClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_):
            return None

    def row(series_id, value=1):
        return {
            "series_id": series_id, "observation_date": "2026-06-30",
            "value": value, "unit": "MW", "source_url": "https://example.test",
        }

    async def monthly(_service, _client):
        return [row("us_electricity_sales_all")]

    async def daily(_service, _client):
        return [row("us_electricity_daily_demand_us48")]

    async def failed_pipeline(_service, _client):
        raise RuntimeError("temporary workbook failure")

    monkeypatch.setattr(settings, "eia_api_key", "test-key")
    monkeypatch.setattr("app.services.regime_eia.httpx.AsyncClient", lambda **_: FakeClient())
    monkeypatch.setattr(EiaPowerService, "_fetch_monthly_context", monthly)
    monkeypatch.setattr(EiaPowerService, "_fetch_daily_grid", daily)
    monkeypatch.setattr(EiaPowerService, "_fetch_capacity_pipeline", failed_pipeline)
    service = EiaPowerService()
    service.repo.save_observations("eia", "us_power", [row("us_power_operating_capacity_mw", 1234)])

    result = await service.refresh(force=True)

    assert result["status"] == "partial"
    assert "pipeline" in result["errors"]
    assert service.repo.series("us_power_operating_capacity_mw")[-1]["value"] == 1234
